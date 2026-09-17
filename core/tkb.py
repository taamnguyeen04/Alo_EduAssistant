"""Đọc thời khóa biểu từ file .xlsx (GV gửi) thành dict lưu vào hồ sơ HS.

Cấu trúc file chuẩn (xem thoi_khoa_bieu_11A2.xlsx):
  - Vài dòng tiêu đề (Lớp, ngày hiệu lực) — parser tự nhặt, thiếu cũng được.
  - Mỗi buổi là 1 bảng: dòng tiêu đề buổi ("Buổi sáng"/"Buổi chiều"),
    dòng header (ô đầu "Tiết", các ô sau "Thứ 2".. "Thứ 7"/"Chủ nhật"),
    rồi các dòng tiết (ô đầu số tiết, các ô sau tên môn, ô trống = không học).
Parser chịu được: sheet tên bất kỳ, header viết không dấu, thứ tự bảng đảo,
dòng trống xen giữa, ô môn xuống dòng (\\n) trong 1 ô.
"""
import io
import re
import unicodedata

NGAY_CHUAN = ["T2", "T3", "T4", "T5", "T6", "T7", "CN"]

_BAN_DAU = {
    "thu 2": "T2", "thu 3": "T3", "thu 4": "T4", "thu 5": "T5",
    "thu 6": "T6", "thu 7": "T7", "thu 8": "CN", "chu nhat": "CN", "cn": "CN",
    "t2": "T2", "t3": "T3", "t4": "T4", "t5": "T5", "t6": "T6", "t7": "T7",
}


def _khong_dau(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def _ma_ngay(o):
    return _BAN_DAU.get(_khong_dau(o), "")


def _la_header_tiet(cells):
    if _khong_dau(cells[0]) not in ("tiet",):
        return False
    return any(_ma_ngay(c) for c in cells[1:])


def _ten_buoi(cells_phia_tren):
    for row in reversed(cells_phia_tren):
        t = _khong_dau(" ".join(str(v or "") for v in row))
        if "sang" in t:
            return "sáng"
        if "chieu" in t:
            return "chiều"
        if "toi" in t:
            return "tối"
    return ""


def _la_dong_ket_thuc(first):
    t = _khong_dau(first)
    return any(k in t for k in ("buoi", "mon chuyen de", "ghi chu", "chu thich"))


def parse_tkb_xlsx(file_bytes, ten_file=""):
    """Parse bytes .xlsx -> {"nguon","lop","hieu_luc","lich"}.

    lich = {T2: {"sáng": {"1": "Vật lí", ...}, "chiều": {...}}, ...}
    Raise ValueError nếu không tìm thấy bảng TKB nào.
    """
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception:
        raise ValueError("File không phải Excel .xlsx hợp lệ.")
    ws = wb[wb.sheetnames[0]]
    rows = [[(str(v).strip() if v is not None else "") for v in r]
            for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    if not rows:
        raise ValueError("File trống, không đọc được thời khóa biểu.")

    # Meta: Lớp / ngày hiệu lực ở vài dòng đầu
    lop, hieu_luc = "", ""
    for r in rows[:4]:
        for i, o in enumerate(r):
            t = _khong_dau(o)
            nxt = r[i + 1] if i + 1 < len(r) else ""
            if t in ("lop",) and nxt:
                lop = nxt
            if any(k in t for k in ("co tac dung", "hieu luc", "ap dung", "tu ngay")) and nxt:
                hieu_luc = nxt

    # Tìm mọi dòng header tiết
    idx_headers = [i for i, r in enumerate(rows) if _la_header_tiet(r)]
    if not idx_headers:
        # Không có bảng Tiết/Thứ thì file sai mẫu, báo để HS chọn file khác.
        raise ValueError("Không tìm thấy bảng nào có cột 'Tiết / Thứ 2..' trong file.")

    lich = {}
    for n, h in enumerate(idx_headers):
        header = rows[h]
        cot_ngay = [(c, _ma_ngay(header[c])) for c in range(1, len(header))]
        cot_ngay = [(c, m) for c, m in cot_ngay if m]
        if not cot_ngay:
            continue
        buoi = _ten_buoi(rows[max(0, h - 5):h]) or ("sáng" if n == 0 else "chiều" if n == 1 else f"buổi {n + 1}")
        ket_thuc = idx_headers[n + 1] if n + 1 < len(idx_headers) else len(rows)
        dem_trong = 0
        for r in rows[h + 1:ket_thuc]:
            o_dau = (r[0] if r else "").strip()
            if not o_dau and not any((c or "").strip() for c in r[1:]):
                dem_trong += 1
                if dem_trong >= 10:
                    break
                continue
            if _la_dong_ket_thuc(o_dau) or _khong_dau(o_dau) == "tiet":
                break
            dem_trong = 0
            m_tiet = re.search(r"\d+", o_dau)
            if not m_tiet:
                continue
            tiet = m_tiet.group(0)
            for c, ma in cot_ngay:
                mon = re.sub(r"\s+", " ", (r[c] if c < len(r) else "")).strip()
                if mon:
                    lich.setdefault(ma, {}).setdefault(buoi, {})[tiet] = mon

    if not lich:
        raise ValueError("Tìm thấy khung bảng nhưng không có tiết học nào.")
    return {"nguon": ten_file, "lop": lop, "hieu_luc": hieu_luc, "lich": lich}


def la_tkb_file(raw):
    """True nếu tkb đã là dạng có cấu trúc từ file (có key 'lich')."""
    return isinstance(raw, dict) and isinstance(raw.get("lich"), dict)


def chuan_hoa_tkb(raw):
    """Chuẩn hóa mọi dạng tkb cũ/mới về {"nguon","lop","hieu_luc","lich"}.

    Dạng cũ {T2: [môn,...]} -> dồn vào buổi "sáng", đánh số tiết 1..n.
    """
    # Đã đúng dạng file thì giữ nguyên để không mất buổi/tiết.
    if la_tkb_file(raw):
        return raw
    lich = {}
    if isinstance(raw, dict):
        for ngay, ds in raw.items():
            if ngay in NGAY_CHUAN and isinstance(ds, list):
                mons = [str(m).strip() for m in ds if str(m).strip()]
                if mons:
                    lich[ngay] = {"sáng": {str(i + 1): m for i, m in enumerate(mons)}}
    return {"nguon": "", "lop": "", "hieu_luc": "", "lich": lich}
